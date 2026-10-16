
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
CSV_FILE = BASE_DIR / "raw_data.csv"
OUTPUT_FILE = BASE_DIR / "sales_dashboard.xlsx"


def main():
    if not CSV_FILE.exists():
        raise FileNotFoundError(f"CSV file not found: {CSV_FILE}")

    df = pd.read_csv(CSV_FILE)
    df.columns = df.columns.str.strip()

    required = ["Order ID", "Order Date", "Category", "Sales"]
    missing = [col for col in required if col not in df.columns]
    if missing:
        raise ValueError(f"Missing required columns: {missing}")

    df["Order Date"] = pd.to_datetime(
        df["Order Date"], errors="coerce"
    )
    df["Sales"] = pd.to_numeric(df["Sales"], errors="coerce")
    df = df.dropna(subset=["Order Date", "Sales", "Category"])
    df = df.drop_duplicates()

    total_sales = df["Sales"].sum()
    unique_orders = df["Order ID"].nunique()

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"

    # Dashboard title
    ws.merge_cells("A1:H1")
    ws["A1"] = "SUPERSTORE SALES DASHBOARD"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 35

    # KPI section
    kpis = [
        ("A3", "Total Sales", "A4", total_sales),
        ("C3", "Unique Orders", "C4", unique_orders),
        ("E3", "Total Profit", "E4", "Not available"),
        ("G3", "Profit Margin", "G4", "Not available"),
    ]

    for label_cell, label, value_cell, value in kpis:
        ws[label_cell] = label
        ws[label_cell].font = Font(bold=True, color="FFFFFF")
        ws[label_cell].fill = PatternFill("solid", fgColor="4472C4")
        ws[value_cell] = value
        ws[value_cell].font = Font(size=14, bold=True)
        ws[value_cell].fill = PatternFill("solid", fgColor="D9EAF7")

    ws["A4"].number_format = '$#,##0.00'
    ws["E4"].comment = __import__("openpyxl").comments.Comment(
        "The source CSV contains no Profit or Cost column.", "ChatGPT"
    )
    ws["G4"].comment = __import__("openpyxl").comments.Comment(
        "Profit margin cannot be calculated without profit data.", "ChatGPT"
    )

    # Sales by category
    category = (
        df.groupby("Category", as_index=False)["Sales"]
        .sum()
        .sort_values("Sales", ascending=False)
    )

    cat_ws = wb.create_sheet("Sales by Category")
    cat_ws.append(["Category", "Total Sales"])
    for row in category.itertuples(index=False, name=None):
        cat_ws.append(list(row))
    for cell in cat_ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for r in range(2, cat_ws.max_row + 1):
        cat_ws.cell(r, 2).number_format = '$#,##0.00'

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Sales by Category"
    chart.x_axis.title = "Sales"
    chart.y_axis.title = "Category"
    chart.add_data(
        Reference(cat_ws, min_col=2, min_row=1, max_row=cat_ws.max_row),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(cat_ws, min_col=1, min_row=2, max_row=cat_ws.max_row)
    )
    chart.height = 7
    chart.width = 12
    ws.add_chart(chart, "A7")

    # Monthly sales
    df["Month"] = df["Order Date"].dt.to_period("M").astype(str)
    monthly = df.groupby("Month", as_index=False)["Sales"].sum()

    month_ws = wb.create_sheet("Monthly Sales")
    month_ws.append(["Month", "Total Sales"])
    for row in monthly.itertuples(index=False, name=None):
        month_ws.append(list(row))
    for cell in month_ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for r in range(2, month_ws.max_row + 1):
        month_ws.cell(r, 2).number_format = '$#,##0.00'

    line = LineChart()
    line.title = "Monthly Sales Trend"
    line.y_axis.title = "Sales"
    line.x_axis.title = "Month"
    line.add_data(
        Reference(month_ws, min_col=2, min_row=1, max_row=month_ws.max_row),
        titles_from_data=True,
    )
    line.set_categories(
        Reference(month_ws, min_col=1, min_row=2, max_row=month_ws.max_row)
    )
    line.height = 7
    line.width = 16
    ws.add_chart(line, "A23")

    # Cleaned data sheet
    data_ws = wb.create_sheet("Cleaned Data")
    data_ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        data_ws.append(list(row))

    # Read Me
    readme = wb.create_sheet("Read Me")
    notes = [
        ["SALES DASHBOARD - NOTES"],
        ["Total Sales is calculated from the Sales column."],
        ["Unique Orders counts distinct Order IDs."],
        ["The source CSV does not contain Profit or Cost data."],
        ["Profit and Profit Margin are therefore not calculated."],
        ["Monthly Sales groups orders by calendar month."],
    ]
    for row in notes:
        readme.append(row)
    readme["A1"].font = Font(bold=True, size=16)

    # Format columns and freeze headers
    for sheet in [cat_ws, month_ws, data_ws]:
        sheet.freeze_panes = "A2"
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            sheet.column_dimensions[letter].width = min(
                max(len(str(sheet.cell(r, col).value or ""))
                    for r in range(1, min(sheet.max_row, 100) + 1)) + 2,
                32,
            )
        sheet.auto_filter.ref = sheet.dimensions

    for sheet in [ws, cat_ws, month_ws, readme]:
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="center", wrap_text=True
                )

    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 18

    wb.save(OUTPUT_FILE)

    print(f"Dashboard created: {OUTPUT_FILE}")
    print(f"Rows in cleaned dataset: {len(df):,}")
    print(f"Total sales: ${total_sales:,.2f}")
    print(f"Unique orders: {unique_orders:,}")


if __name__ == "__main__":
    main()
